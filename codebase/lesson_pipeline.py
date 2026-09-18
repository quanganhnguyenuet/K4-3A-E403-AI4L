"""Create reviewed, source-grounded lessons from uploaded PDF or PPTX files.

Pipeline: Document Parser -> Extraction Agent -> Validation/Critic Agent
(refinement loop) -> Lesson Generator -> human review -> publish. The
extraction agent builds a major/sub concept knowledge structure from the
whole document (not slide-by-slide); the critic checks that structure before
any teach-back content is written from it.
"""
from __future__ import annotations

import base64
import io
import json
import os
import re
import sqlite3
import statistics
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from agent_core import AuditLogger, OpenAIResponsesProvider
from platform_runtime import LessonCatalog

MAX_FILE_BYTES = 12_000_000
MAX_SEGMENTS = 160
MAX_SOURCE_TEXT = 180_000
MAX_VALIDATION_ROUNDS = 2
DEFAULT_DRAFT_DB = Path(__file__).resolve().parent / "state" / "lesson_drafts.sqlite3"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _parse_pdf(content: bytes) -> list[dict[str, Any]]:
    import pymupdf

    document = pymupdf.open(stream=content, filetype="pdf")
    if document.page_count > 200:
        raise ValueError("PDF vượt giới hạn 200 trang")
    units: list[dict[str, Any]] = []
    for page_no, page in enumerate(document, 1):
        raw_blocks = [b for b in page.get_text("dict").get("blocks", []) if b.get("type") == 0]
        raw_blocks.sort(key=lambda b: (round(b["bbox"][1], 1), round(b["bbox"][0], 1)))
        sizes = [span["size"] for b in raw_blocks for line in b.get("lines", []) for span in line.get("spans", [])]
        median_size = statistics.median(sizes) if sizes else 0.0
        title_index, best_size = None, 0.0
        for idx, block in enumerate(raw_blocks):
            block_sizes = [span["size"] for line in block.get("lines", []) for span in line.get("spans", [])]
            if block_sizes and max(block_sizes) > median_size * 1.15 and max(block_sizes) > best_size:
                best_size, title_index = max(block_sizes), idx
        segments, title_text = [], None
        for block_no, block in enumerate(raw_blocks, 1):
            text = _clean(" ".join(span["text"] for line in block.get("lines", []) for span in line.get("spans", [])))
            if len(text) < 20:
                continue
            kind = "title" if block_no - 1 == title_index else "body"
            segments.append({
                "id": f"SRC-{page_no:03d}-{block_no:03d}", "kind": kind, "text": text,
                "locator": f"trang {page_no}, khối {block_no}",
            })
            if kind == "title" and title_text is None:
                title_text = text
        units.append({"unit_id": f"PAGE-{page_no:03d}", "index": page_no, "title": title_text, "segments": segments})
    return units


def _parse_pptx(content: bytes) -> list[dict[str, Any]]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    if len(presentation.slides) > 300:
        raise ValueError("PPTX vượt giới hạn 300 slide")
    units: list[dict[str, Any]] = []
    for slide_no, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        segments, title_text = [], None
        for shape_no, shape in enumerate(slide.shapes, 1):
            parts = []
            if shape.has_text_frame:
                parts.extend(p.text for p in shape.text_frame.paragraphs)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            text = _clean(" ".join(parts))
            if len(text) < 20:
                continue
            if shape.has_table:
                kind = "table"
            elif title_shape is not None and shape.shape_id == title_shape.shape_id:
                kind = "title"
            else:
                kind = "body"
            segments.append({
                "id": f"SRC-{slide_no:03d}-{shape_no:03d}", "kind": kind, "text": text,
                "locator": f"slide {slide_no}, đối tượng {shape_no}",
            })
            if kind == "title" and title_text is None:
                title_text = text
        units.append({"unit_id": f"SLIDE-{slide_no:03d}", "index": slide_no, "title": title_text, "segments": segments})
    return units


def extract_document(filename: str, content: bytes) -> dict[str, Any]:
    """Return a flat evidence-citable segment list plus the ordered page/slide grouping.

    ``segments`` keeps the exact shape earlier callers rely on (id/file/type/locator/text),
    since ``validate_draft`` and published lessons key evidence off ``segment["id"]``.
    ``units`` preserves document order and slide/page boundaries for the extraction agent,
    so it sees the lecture's real structure instead of reconstructing it from locator strings.
    """
    name = Path(filename).name
    if not name or not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Tệp rỗng hoặc vượt giới hạn 12 MB")
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            units = _parse_pdf(content)
        elif suffix == ".pptx":
            units = _parse_pptx(content)
        else:
            raise ValueError("Chỉ hỗ trợ tệp .pdf hoặc .pptx")
    except ImportError as exc:
        raise RuntimeError("Thiếu thư viện đọc tài liệu; hãy cài requirements.txt") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Không đọc được tài liệu: {exc}") from exc
    file_type = suffix.lstrip(".")
    segments = [
        {"id": seg["id"], "file": name, "type": file_type, "locator": seg["locator"], "text": seg["text"]}
        for unit in units
        for seg in unit["segments"]
    ]
    if not segments:
        raise ValueError("Không trích được chữ; tài liệu scan hoặc slide ảnh cần OCR trước")
    if len(segments) > MAX_SEGMENTS or sum(len(s["text"]) for s in segments) > MAX_SOURCE_TEXT:
        raise ValueError("Tài liệu quá dài; hãy tách thành bài nhỏ hơn trước khi tạo lesson")
    return {
        "segments": segments,
        "units": [
            {
                "unit_id": unit["unit_id"], "index": unit["index"], "title": unit["title"],
                "segments": [{"id": s["id"], "kind": s["kind"], "text": s["text"]} for s in unit["segments"]],
            }
            for unit in units
        ],
    }


def _schema(source_ids: list[str]) -> dict[str, Any]:
    string = {"type": "string"}; strings = {"type": "array", "items": string}
    evidence = {"type": "object", "additionalProperties": False, "required": ["source_id", "quote"], "properties": {"source_id": {"type": "string", "enum": source_ids}, "quote": string}}
    point = {"type": "object", "additionalProperties": False, "required": ["id", "label", "weight", "ground_truth", "sub_concepts", "accepted_signals", "signals", "question", "recovery", "support_evidence"], "properties": {"id": string, "label": string, "weight": {"type": "integer"}, "ground_truth": string, "sub_concepts": strings, "accepted_signals": strings, "signals": strings, "question": string, "recovery": string, "support_evidence": {"type": "array", "items": evidence}}}
    misconception = {"type": "object", "additionalProperties": False, "required": ["id", "claim", "explanation", "socratic_question", "conflicts_with", "signals"], "properties": {"id": string, "claim": string, "explanation": string, "socratic_question": string, "conflicts_with": strings, "signals": strings}}
    lesson = {"type": "object", "additionalProperties": False, "required": ["id", "track", "title", "description", "greeting", "task", "scope_terms", "starter_prompts", "transfer_question", "points", "misconceptions"], "properties": {"id": string, "track": string, "title": string, "description": string, "greeting": string, "task": string, "scope_terms": strings, "starter_prompts": strings, "transfer_question": string, "points": {"type": "array", "items": point}, "misconceptions": {"type": "array", "items": misconception}}}
    return {"type": "object", "additionalProperties": False, "required": ["lesson"], "properties": {"lesson": lesson}}


def _structure_schema(source_ids: list[str]) -> dict[str, Any]:
    string = {"type": "string"}; strings = {"type": "array", "items": string}
    evidence = {"type": "object", "additionalProperties": False, "required": ["source_id", "quote"], "properties": {"source_id": {"type": "string", "enum": source_ids}, "quote": string}}
    evidences = {"type": "array", "items": evidence}
    sub_concept = {"type": "object", "additionalProperties": False, "required": ["id", "name", "kind", "evidence"], "properties": {"id": string, "name": string, "kind": {"type": "string", "enum": ["definition", "property", "example", "process", "comparison", "other"]}, "evidence": evidences}}
    major_concept = {"type": "object", "additionalProperties": False, "required": ["id", "name", "summary", "evidence", "sub_concepts"], "properties": {"id": string, "name": string, "summary": string, "evidence": evidences, "sub_concepts": {"type": "array", "items": sub_concept}}}
    return {"type": "object", "additionalProperties": False, "required": ["lesson_title", "learning_objectives", "major_concepts"], "properties": {"lesson_title": string, "learning_objectives": strings, "major_concepts": {"type": "array", "items": major_concept}}}


def _critic_schema() -> dict[str, Any]:
    string = {"type": "string"}
    issue = {"type": "object", "additionalProperties": False, "required": ["type", "concept_id", "detail", "suggestion"], "properties": {"type": {"type": "string", "enum": ["missing_concept", "duplicate_concept", "wrong_granularity", "misplaced_subconcept", "example_as_concept", "metadata_as_concept", "unsupported_claim", "order_violation"]}, "concept_id": {"type": ["string", "null"]}, "detail": string, "suggestion": string}}
    return {"type": "object", "additionalProperties": False, "required": ["verdict", "issues"], "properties": {"verdict": {"type": "string", "enum": ["approved", "needs_revision"]}, "issues": {"type": "array", "items": issue}}}


def extract_knowledge_structure(units: list[dict[str, Any]], segments: list[dict[str, Any]], provider: OpenAIResponsesProvider, *, session_id: str) -> dict[str, Any]:
    """Extraction Agent: build a major/sub concept knowledge structure from the whole document."""
    return provider._request_structured(
        instructions=(
            "Bạn là một Extraction Agent xây dựng cấu trúc kiến thức của một bài giảng, không phải tóm tắt tài liệu. "
            "Đầu vào là các slide/trang (units) theo đúng thứ tự trình bày, mỗi segment có kind title/body/table. "
            "Nhiệm vụ: xác định Major Concept (kiến thức/chủ đề lớn người học phải nắm được) và Sub Concept (định "
            "nghĩa, tính chất, ví dụ, quy trình... giúp hiểu Major Concept đó). TUYỆT ĐỐI không biến mỗi slide hay "
            "mỗi bullet thành một Major Concept riêng; hãy tổng hợp trên toàn bộ tài liệu trước khi nhóm. Số lượng "
            "Major Concept phụ thuộc nội dung thực tế, không cố định và không cần chia đều. Không nhầm heading/"
            "metadata (tên tác giả, số trang, mục lục) hoặc ví dụ minh hoạ thành kiến thức cốt lõi. Mỗi concept và "
            "sub-concept phải có ít nhất một evidence với source_id có thật và quote chép NGUYÊN VĂN từ segment "
            "tương ứng; không bịa dữ kiện. Giữ đúng thứ tự logic của bài giảng."
        ),
        input_text=json.dumps({"units": units}, ensure_ascii=False),
        schema=_structure_schema([s["id"] for s in segments]),
        schema_name="lesson_knowledge_structure",
        component="lesson_extraction_agent",
        session_id=session_id,
        max_output_tokens=4000,
    )


def refine_knowledge_structure(units: list[dict[str, Any]], segments: list[dict[str, Any]], structure: dict[str, Any], issues: list[dict[str, Any]], provider: OpenAIResponsesProvider, *, session_id: str) -> dict[str, Any]:
    """Extraction Agent, second pass: fix only the issues the critic raised."""
    return provider._request_structured(
        instructions=(
            "Bạn là Extraction Agent đang SỬA LẠI cấu trúc kiến thức đã tạo trước đó theo phản hồi của Validation "
            "Agent. Chỉ sửa đúng những vấn đề trong issues_to_fix (gộp concept trùng lặp, tách lại granularity, "
            "chuyển sub-concept về đúng concept cha, bỏ ví dụ/metadata bị nhầm thành concept, bổ sung concept còn "
            "thiếu, gắn lại evidence cho đúng...), giữ nguyên phần đã đúng trong previous_structure. Vẫn phải bám "
            "sát toàn bộ nội dung units gốc và chỉ dùng evidence có source_id thật với quote nguyên văn."
        ),
        input_text=json.dumps({"units": units, "previous_structure": structure, "issues_to_fix": issues}, ensure_ascii=False),
        schema=_structure_schema([s["id"] for s in segments]),
        schema_name="lesson_knowledge_structure",
        component="lesson_extraction_agent",
        session_id=session_id,
        max_output_tokens=4000,
    )


def critique_structure(structure: dict[str, Any], segments: list[dict[str, Any]], provider: OpenAIResponsesProvider, *, session_id: str) -> dict[str, Any]:
    """Validation Agent / critic: check the knowledge structure against the source segments."""
    return provider._request_structured(
        instructions=(
            "Bạn là Validation Agent / LLM Critic kiểm tra cấu trúc kiến thức của một bài giảng. Dựa trên các đoạn "
            "nguồn gốc (segments) và cấu trúc major/sub concept đã trích xuất, hãy kiểm tra: có bỏ sót concept quan "
            "trọng không; có concept trùng nhau không; mỗi Major Concept có thực sự là kiến thức lớn (không phải "
            "một chi tiết nhỏ) không; sub-concept có thuộc đúng concept cha không; có ví dụ nào bị nhầm thành kiến "
            "thức cốt lõi không; có heading/metadata bị nhầm thành kiến thức không; mọi evidence có đúng và có thật "
            "trong segments không; cấu trúc có giữ đúng thứ tự logic của bài giảng không. Trả verdict approved chỉ "
            "khi không còn vấn đề đáng kể; ngược lại trả needs_revision kèm issues cụ thể, mỗi issue nêu rõ type, "
            "concept_id liên quan (hoặc null nếu là concept còn thiếu), detail và suggestion để sửa."
        ),
        input_text=json.dumps({"segments": [{"id": s["id"], "text": s["text"]} for s in segments], "structure": structure}, ensure_ascii=False),
        schema=_critic_schema(),
        schema_name="lesson_structure_critique",
        component="lesson_validation_agent",
        session_id=session_id,
        max_output_tokens=2000,
    )


def build_validated_structure(units: list[dict[str, Any]], segments: list[dict[str, Any]], provider: OpenAIResponsesProvider, *, session_id: str, max_rounds: int = MAX_VALIDATION_ROUNDS) -> tuple[dict[str, Any], list[str]]:
    """Run Extraction -> Critic, refining up to ``max_rounds`` times. Returns (structure, unresolved_issues)."""
    structure = extract_knowledge_structure(units, segments, provider, session_id=session_id)
    critique = critique_structure(structure, segments, provider, session_id=session_id)
    rounds = 0
    while critique.get("verdict") != "approved" and critique.get("issues") and rounds < max_rounds:
        structure = refine_knowledge_structure(units, segments, structure, critique["issues"], provider, session_id=session_id)
        critique = critique_structure(structure, segments, provider, session_id=session_id)
        rounds += 1
    unresolved = [] if critique.get("verdict") == "approved" else [
        f"{issue.get('type')}: {issue.get('detail')}" for issue in critique.get("issues", [])
    ]
    return structure, unresolved


def generate_lesson_from_structure(structure: dict[str, Any], segments: list[dict[str, Any]], provider: OpenAIResponsesProvider, *, session_id: str) -> dict[str, Any]:
    """Lesson Generator: turn a validated knowledge structure into the teach-back lesson schema."""
    return provider._request_structured(
        instructions=(
            "Bạn là Lesson Generator. Đầu vào là một cấu trúc kiến thức (major_concepts/sub_concepts) đã được trích "
            "xuất và kiểm định từ một bài giảng — KHÔNG tự ý thêm, bớt hay gộp concept nữa. Với mỗi major_concept, "
            "tạo đúng một point teach-back bằng tiếng Việt: label lấy từ name, sub_concepts là danh sách tên các "
            "sub_concept của concept đó, ground_truth là một quote tiêu biểu trong evidence, question và recovery "
            "viết tự nhiên, accepted_signals/signals bám sát nội dung concept, support_evidence copy nguyên văn "
            "evidence của concept đó (source_id thật, quote nguyên văn). Không bịa dữ kiện ngoài structure và "
            "segments đã cho. Weight của các point cộng thành 100; chia đều trừ khi vai trò học tập thực sự khác "
            "nhau rõ rệt. Dùng point id P1..Pn, misconception id C1..Cn."
        ),
        input_text=json.dumps({"structure": structure, "segments": [{"id": s["id"], "text": s["text"]} for s in segments]}, ensure_ascii=False),
        schema=_schema([s["id"] for s in segments]),
        schema_name="lesson_draft",
        component="lesson_generator",
        session_id=session_id,
        max_output_tokens=3200,
    )


def repair_evidence_quotes(lesson: dict[str, Any], segments: list[dict[str, Any]]) -> None:
    """The model selects source IDs semantically; preserve exact source text for validation."""
    source_by_id = {segment["id"]: segment["text"] for segment in segments}
    for point in lesson.get("points", []):
        for evidence in point.get("support_evidence", []):
            source = source_by_id.get(evidence.get("source_id"), "")
            quote = str(evidence.get("quote", "")).strip()
            if source and (len(quote) < 12 or quote not in source):
                evidence["quote"] = source[: min(280, len(source))]


def validate_draft(generated: dict[str, Any], segments: list[dict[str, str]], existing_ids: set[str]) -> tuple[dict[str, Any] | None, list[str]]:
    """Accept only structurally-valid lessons whose quoted evidence occurs in the upload."""
    errors: list[str] = []
    try:
        raw = generated["lesson"]
        if not re.fullmatch(r"[a-z][a-z0-9-]{2,63}", raw["id"]): errors.append("ID lesson phải là slug chữ thường")
        if raw["id"] in existing_ids: errors.append("ID lesson đã tồn tại")
        points = raw["points"]
        if not 2 <= len(points) <= 6: errors.append("Bài cần từ 2 đến 6 ý chính")
        if sum(int(p["weight"]) for p in points) != 100: errors.append("Trọng số các ý phải cộng thành 100")
        source_map = {s["id"]: s for s in segments}
        source_ids: set[str] = set()
        for point in points:
            if not point["support_evidence"]: errors.append(f"{point['id']}: thiếu dẫn chứng nguồn")
            for evidence in point["support_evidence"]:
                source = source_map.get(evidence["source_id"]); quote = evidence["quote"].strip()
                if not source or len(quote) < 12 or quote not in source["text"]:
                    errors.append(f"{point['id']}: trích dẫn không nằm nguyên văn trong bài giảng")
                else: source_ids.add(source["id"])
            point["source_ids"] = [e["source_id"] for e in point["support_evidence"] if e["source_id"] in source_map]
            point["required"] = True
        raw["sources"] = [{"id": s["id"], "type": s["type"], "file": s["file"], "locator": s["locator"], "paraphrase": s["text"][:240], "quote": s["text"][:360]} for s in segments if s["id"] in source_ids]
        raw["transfer_point_ids"] = [p["id"] for p in points]
        raw["mastery_criteria"] = {"required_point_ids": [p["id"] for p in points], "require_no_misconceptions": True, "require_transfer_example": True}
        LessonCatalog._validate_lesson(raw)
        return (None, errors) if errors else (raw, [])
    except (KeyError, TypeError, ValueError) as exc:
        return None, [f"Cấu trúc bản nháp không hợp lệ: {exc}"]


class LessonDraftStore:
    def __init__(self, path: Path = DEFAULT_DRAFT_DB) -> None:
        self.path = Path(path); self.path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(path), check_same_thread=False); self.conn.row_factory = sqlite3.Row; self.lock = threading.RLock()
        self.conn.execute(
            "CREATE TABLE IF NOT EXISTS lesson_drafts ("
            "id TEXT PRIMARY KEY, filename TEXT NOT NULL, status TEXT NOT NULL, "
            "segments_json TEXT NOT NULL, generated_json TEXT NOT NULL, lesson_json TEXT, "
            "errors_json TEXT NOT NULL, structure_json TEXT NOT NULL DEFAULT '{}', "
            "critic_json TEXT NOT NULL DEFAULT '[]', created_at TEXT NOT NULL, updated_at TEXT NOT NULL)"
        )
        existing = {row["name"] for row in self.conn.execute("PRAGMA table_info(lesson_drafts)").fetchall()}
        migrations = {
            "structure_json": "ALTER TABLE lesson_drafts ADD COLUMN structure_json TEXT NOT NULL DEFAULT '{}'",
            "critic_json": "ALTER TABLE lesson_drafts ADD COLUMN critic_json TEXT NOT NULL DEFAULT '[]'",
        }
        for name, statement in migrations.items():
            if name not in existing:
                self.conn.execute(statement)
        self.conn.commit()

    def save(self, filename: str, segments: list[dict[str, str]], generated: dict[str, Any], lesson: dict[str, Any] | None, errors: list[str], draft_id: str | None = None, structure: dict[str, Any] | None = None, critic_issues: list[str] | None = None) -> dict[str, Any]:
        with self.lock:
            identifier, now = draft_id or str(uuid.uuid4()), _now(); status = "needs_changes" if errors else "ready_for_review"
            values = (
                filename, status, json.dumps(segments, ensure_ascii=False), json.dumps(generated, ensure_ascii=False),
                json.dumps(lesson, ensure_ascii=False), json.dumps(errors, ensure_ascii=False),
                json.dumps(structure or {}, ensure_ascii=False), json.dumps(critic_issues or [], ensure_ascii=False), now,
            )
            if draft_id:
                self.conn.execute(
                    "UPDATE lesson_drafts SET filename=?, status=?, segments_json=?, generated_json=?, lesson_json=?, errors_json=?, structure_json=?, critic_json=?, updated_at=? WHERE id=?",
                    (*values, identifier),
                )
            else:
                self.conn.execute("INSERT INTO lesson_drafts VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", (identifier, *values, now))
            self.conn.commit(); return self.get(identifier)

    def get(self, draft_id: str) -> dict[str, Any]:
        row = self.conn.execute("SELECT * FROM lesson_drafts WHERE id=?", (draft_id,)).fetchone()
        if not row: raise KeyError("Không tìm thấy bản nháp")
        return {
            "id": row["id"], "filename": row["filename"], "status": row["status"],
            "segments": json.loads(row["segments_json"]), "generated": json.loads(row["generated_json"]),
            "lesson": json.loads(row["lesson_json"]), "errors": json.loads(row["errors_json"]),
            "structure": json.loads(row["structure_json"]), "critic_issues": json.loads(row["critic_json"]),
            "created_at": row["created_at"], "updated_at": row["updated_at"],
        }

    def mark_published(self, draft_id: str) -> None:
        self.conn.execute("UPDATE lesson_drafts SET status='published', updated_at=? WHERE id=?", (_now(), draft_id)); self.conn.commit()


class LessonDraftService:
    def __init__(self, catalog: LessonCatalog, store: LessonDraftStore | None = None, on_publish: Any = None, logger: AuditLogger | None = None) -> None:
        self.catalog, self.store, self.on_publish, self.logger = catalog, store or LessonDraftStore(), on_publish, logger or AuditLogger(); self.lock = threading.RLock()

    def create(self, filename: str, content_base64: str, api_key: str | None, model: str | None = None) -> dict[str, Any]:
        try: content = base64.b64decode(content_base64, validate=True)
        except Exception as exc: raise ValueError("Nội dung tệp không phải base64 hợp lệ") from exc
        parsed = extract_document(filename, content)
        segments, units = parsed["segments"], parsed["units"]
        provider = OpenAIResponsesProvider(api_key=api_key or os.getenv("OPENAI_API_KEY"), model=model, logger=self.logger)
        session_id = f"lesson-draft-{uuid.uuid4()}"
        structure, unresolved_issues = build_validated_structure(units, segments, provider, session_id=session_id)
        generated = generate_lesson_from_structure(structure, segments, provider, session_id=session_id)
        if isinstance(generated.get("lesson"), dict):
            repair_evidence_quotes(generated["lesson"], segments)
        lesson, errors = validate_draft(generated, segments, set(self.catalog._lessons))
        errors = errors + unresolved_issues
        return self.store.save(Path(filename).name, segments, generated, lesson, errors, structure=structure, critic_issues=unresolved_issues)

    def revise(self, draft_id: str, generated: dict[str, Any]) -> dict[str, Any]:
        draft = self.store.get(draft_id)
        if isinstance(generated.get("lesson"), dict):
            repair_evidence_quotes(generated["lesson"], draft["segments"])
        lesson, errors = validate_draft(generated, draft["segments"], set(self.catalog._lessons))
        return self.store.save(draft["filename"], draft["segments"], generated, lesson, errors, draft_id, structure=draft["structure"], critic_issues=draft["critic_issues"])

    def publish(self, draft_id: str) -> dict[str, Any]:
        with self.lock:
            draft = self.store.get(draft_id)
            if draft["status"] != "ready_for_review": raise ValueError("Bản nháp chưa đạt kiểm định")
            lesson, errors = validate_draft(draft["generated"], draft["segments"], set(self.catalog._lessons))
            if errors or not lesson: raise ValueError("Bản nháp không còn hợp lệ: " + "; ".join(errors))
            payload = json.loads(self.catalog.path.read_text(encoding="utf-8")); payload["lessons"].append(lesson)
            temp = self.catalog.path.with_suffix(".tmp")
            temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            os.replace(temp, self.catalog.path); self.catalog = LessonCatalog(self.catalog.path)
            if self.on_publish: self.on_publish(self.catalog)
            self.store.mark_published(draft_id); return lesson
