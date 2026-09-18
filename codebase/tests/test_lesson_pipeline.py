from __future__ import annotations

import base64
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

CODEBASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(CODEBASE))

import lesson_pipeline as lp  # noqa: E402
from agent_core import AuditLogger, OpenAIResponsesProvider  # noqa: E402
from platform_runtime import LessonCatalog  # noqa: E402

TITLE_1 = "Introduction to Large Language Models"
TITLE_2 = "How LLM Generates Text End to End"
BODY_1 = "What is an LLM? Next-token prediction basics for language modeling."
BODY_2 = "Tokenization splits text; probability distribution picks the next token."

SEED_CATALOG = {
    "schema_version": "2.0",
    "lessons": [
        {
            "id": "seed-lesson", "track": "Day 0", "title": "Seed", "description": "d",
            "greeting": "g", "task": "t",
            "points": [{"id": "P1", "label": "L", "weight": 100, "source_ids": []}],
            "sources": [{"id": "SRC-000-000", "type": "pptx", "file": "seed.pptx", "locator": "slide 1"}],
        }
    ],
}


def _build_pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide1 = prs.slides.add_slide(layout)
    slide1.shapes.title.text = TITLE_1
    slide1.placeholders[1].text_frame.text = BODY_1
    slide2 = prs.slides.add_slide(layout)
    slide2.shapes.title.text = TITLE_2
    slide2.placeholders[1].text_frame.text = BODY_2
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class ExtractDocumentTests(unittest.TestCase):
    def test_pptx_groups_into_ordered_units_with_titles(self) -> None:
        parsed = lp.extract_document("lecture.pptx", _build_pptx_bytes())
        self.assertEqual(len(parsed["units"]), 2)
        self.assertEqual(parsed["units"][0]["title"], TITLE_1)
        self.assertEqual(parsed["units"][1]["title"], TITLE_2)
        flat_ids = {s["id"] for s in parsed["segments"]}
        unit_ids = {s["id"] for unit in parsed["units"] for s in unit["segments"]}
        self.assertEqual(flat_ids, unit_ids)
        title_segments = [s for unit in parsed["units"] for s in unit["segments"] if s["kind"] == "title"]
        self.assertEqual(len(title_segments), 2)


class BuildValidatedStructureTests(unittest.TestCase):
    def test_refines_once_then_stops_when_critic_approves(self) -> None:
        provider = OpenAIResponsesProvider(api_key="sk-test", model="test-model")
        calls = {"extract": 0, "refine": 0, "critique": 0}
        segments = [{"id": "SRC-001-001", "file": "f.pptx", "type": "pptx", "locator": "slide 1", "text": TITLE_1}]
        units = [{"unit_id": "SLIDE-001", "index": 1, "title": TITLE_1,
                  "segments": [{"id": "SRC-001-001", "kind": "title", "text": TITLE_1}]}]
        structure_v1 = {"lesson_title": "L", "learning_objectives": [], "major_concepts": [
            {"id": "K1", "name": "X", "summary": "s",
             "evidence": [{"source_id": "SRC-001-001", "quote": TITLE_1}], "sub_concepts": []}]}
        structure_v2 = json.loads(json.dumps(structure_v1))
        structure_v2["major_concepts"][0]["name"] = "X fixed"

        def fake_request(**kwargs):
            component = kwargs["component"]
            if component == "lesson_extraction_agent":
                if calls["refine"] == 0 and "issues_to_fix" not in kwargs["input_text"]:
                    calls["extract"] += 1
                    return structure_v1
                calls["refine"] += 1
                self.assertIn("issues_to_fix", kwargs["input_text"])
                return structure_v2
            if component == "lesson_validation_agent":
                calls["critique"] += 1
                if calls["critique"] == 1:
                    return {"verdict": "needs_revision", "issues": [
                        {"type": "wrong_granularity", "concept_id": "K1", "detail": "too narrow", "suggestion": "rename"}
                    ]}
                return {"verdict": "approved", "issues": []}
            raise AssertionError(f"unexpected component {component}")

        provider._request_structured = fake_request  # type: ignore[method-assign]
        structure, unresolved = lp.build_validated_structure(units, segments, provider, session_id="test")
        self.assertEqual(structure["major_concepts"][0]["name"], "X fixed")
        self.assertEqual(unresolved, [])
        self.assertEqual(calls, {"extract": 1, "refine": 1, "critique": 2})

    def test_reports_unresolved_issues_after_max_rounds(self) -> None:
        provider = OpenAIResponsesProvider(api_key="sk-test", model="test-model")
        segments = [{"id": "SRC-001-001", "file": "f.pptx", "type": "pptx", "locator": "slide 1", "text": TITLE_1}]
        units = [{"unit_id": "SLIDE-001", "index": 1, "title": TITLE_1,
                  "segments": [{"id": "SRC-001-001", "kind": "title", "text": TITLE_1}]}]
        structure = {"lesson_title": "L", "learning_objectives": [], "major_concepts": [
            {"id": "K1", "name": "X", "summary": "s",
             "evidence": [{"source_id": "SRC-001-001", "quote": TITLE_1}], "sub_concepts": []}]}
        issue = {"type": "duplicate_concept", "concept_id": "K1", "detail": "still wrong", "suggestion": "merge"}

        def fake_request(**kwargs):
            component = kwargs["component"]
            if component == "lesson_extraction_agent":
                return structure
            if component == "lesson_validation_agent":
                return {"verdict": "needs_revision", "issues": [issue]}
            raise AssertionError(f"unexpected component {component}")

        provider._request_structured = fake_request  # type: ignore[method-assign]
        _, unresolved = lp.build_validated_structure(units, segments, provider, session_id="test", max_rounds=2)
        self.assertEqual(unresolved, ["duplicate_concept: still wrong"])


class GenerateLessonFromStructureTests(unittest.TestCase):
    def test_output_passes_validate_draft_with_concept_driven_point_count(self) -> None:
        provider = OpenAIResponsesProvider(api_key="sk-test", model="test-model")
        segments = [
            {"id": "SRC-001-001", "file": "f.pptx", "type": "pptx", "locator": "slide 1", "text": BODY_1},
            {"id": "SRC-002-001", "file": "f.pptx", "type": "pptx", "locator": "slide 2", "text": BODY_2},
        ]
        structure = {"lesson_title": "L", "learning_objectives": ["obj"], "major_concepts": [
            {"id": "K1", "name": "X", "summary": "s", "evidence": [{"source_id": "SRC-001-001", "quote": BODY_1}], "sub_concepts": []},
            {"id": "K2", "name": "Y", "summary": "s", "evidence": [{"source_id": "SRC-002-001", "quote": BODY_2}], "sub_concepts": []},
        ]}

        def fake_request(**kwargs):
            return {"lesson": {
                "id": "demo-lesson", "track": "Day 1", "title": "L", "description": "d", "greeting": "g",
                "task": "t", "scope_terms": ["X"], "starter_prompts": ["p"], "transfer_question": "q?",
                "points": [
                    {"id": "P1", "label": "X", "weight": 50, "ground_truth": BODY_1, "sub_concepts": [],
                     "accepted_signals": ["a"], "signals": ["s"], "question": "q?", "recovery": "r",
                     "support_evidence": [{"source_id": "SRC-001-001", "quote": BODY_1}]},
                    {"id": "P2", "label": "Y", "weight": 50, "ground_truth": BODY_2, "sub_concepts": [],
                     "accepted_signals": ["a"], "signals": ["s"], "question": "q?", "recovery": "r",
                     "support_evidence": [{"source_id": "SRC-002-001", "quote": BODY_2}]},
                ],
                "misconceptions": [],
            }}

        provider._request_structured = fake_request  # type: ignore[method-assign]
        generated = lp.generate_lesson_from_structure(structure, segments, provider, session_id="test")
        lesson, errors = lp.validate_draft(generated, segments, existing_ids=set())
        self.assertEqual(errors, [])
        self.assertIsNotNone(lesson)
        self.assertEqual(len(lesson["points"]), 2)


class LessonDraftServiceEndToEndTests(unittest.TestCase):
    def _make_service(self, tmp_dir: Path) -> lp.LessonDraftService:
        catalog_path = tmp_dir / "catalog.json"
        catalog_path.write_text(json.dumps(SEED_CATALOG, ensure_ascii=False), encoding="utf-8")
        catalog = LessonCatalog(catalog_path)
        store = lp.LessonDraftStore(tmp_dir / "drafts.sqlite3")
        logger = AuditLogger(tmp_dir / "audit.jsonl")
        return lp.LessonDraftService(catalog, store=store, logger=logger)

    def test_point_count_follows_actual_major_concepts_not_a_fixed_number(self) -> None:
        tmp_dir = Path(tempfile.mkdtemp(prefix="d3-lesson-pipeline-test-"))
        service = self._make_service(tmp_dir)
        major_concepts = [
            {"id": f"K{i}", "name": f"Concept {i}", "summary": "s",
             "evidence": [{"source_id": "SRC-001-001", "quote": TITLE_1}], "sub_concepts": []}
            for i in range(1, 4)
        ]
        structure = {"lesson_title": "L", "learning_objectives": ["o"], "major_concepts": major_concepts}
        weights = [33, 33, 34]

        class FakeProvider:
            def __init__(self, **kwargs) -> None:
                pass

            def _request_structured(self, **kwargs):
                component = kwargs["component"]
                if component == "lesson_extraction_agent":
                    return structure
                if component == "lesson_validation_agent":
                    return {"verdict": "approved", "issues": []}
                if component == "lesson_generator":
                    return {"lesson": {
                        "id": "generated-lesson", "track": "Day 1", "title": "L", "description": "d",
                        "greeting": "g", "task": "t", "scope_terms": [], "starter_prompts": [],
                        "transfer_question": "q?", "misconceptions": [],
                        "points": [
                            {"id": f"P{i + 1}", "label": concept["name"], "weight": weights[i],
                             "ground_truth": TITLE_1, "sub_concepts": [], "accepted_signals": ["a"], "signals": ["s"],
                             "question": "q?", "recovery": "r",
                             "support_evidence": [{"source_id": "SRC-001-001", "quote": TITLE_1}]}
                            for i, concept in enumerate(major_concepts)
                        ],
                    }}
                raise AssertionError(f"unexpected component {component}")

        with mock.patch.object(lp, "OpenAIResponsesProvider", FakeProvider):
            content_b64 = base64.b64encode(_build_pptx_bytes()).decode("ascii")
            draft = service.create("lecture.pptx", content_b64, api_key="sk-test")

        self.assertEqual(draft["status"], "ready_for_review")
        self.assertEqual(len(draft["lesson"]["points"]), 3)
        self.assertEqual(draft["structure"]["major_concepts"][0]["name"], "Concept 1")
        self.assertEqual(draft["critic_issues"], [])


if __name__ == "__main__":
    unittest.main()
